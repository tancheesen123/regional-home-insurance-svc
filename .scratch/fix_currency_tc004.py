# -*- coding: utf-8 -*-
import docx
from docx.oxml.ns import qn
from docx.enum.text import WD_COLOR_INDEX

REPLACEMENTS = [
    ("$500.00", "Rp500"),
    ("$715.00", "Rp715"),
]

def fix_doc(path, save_path=None):
    d = docx.Document(path)
    fixed = 0
    for t in d.tables:
        if len(t.rows[0].cells) <= 1:
            continue
        tc_id = t.rows[0].cells[1].text.strip()
        if tc_id not in ("TC004_01", "TC004_03"):
            continue
        for row in t.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for r in para.runs:
                        for t_el in r._r.findall(qn('w:t')):
                            if not t_el.text:
                                continue
                            new_text = t_el.text
                            for old, new in REPLACEMENTS:
                                if old in new_text:
                                    new_text = new_text.replace(old, new)
                            if new_text != t_el.text:
                                t_el.text = new_text
                                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                                fixed += 1
    out = save_path or path
    try:
        d.save(out)
        print(f"{path}: fixed {fixed} occurrences -> saved to {out}")
    except PermissionError:
        alt = out.replace(".docx", "_LOCKED_RETRY.docx")
        d.save(alt)
        print(f"{path}: fixed {fixed} occurrences -> LOCKED, saved to {alt}")

fix_doc(r"D:\A.UTM Degree\Year-2\Sem 6\PSM2\STD\04 Template STD-PSM1 2024 v2 - updated v7_fix.docx")
fix_doc(r"D:\A.UTM Degree\Year-2\Sem 6\PSM2\Template-Tesis-UTM-v2-PSM-UG-SC-System-Development-u_REVIEWED_v6.docx")
